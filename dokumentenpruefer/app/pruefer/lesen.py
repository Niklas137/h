"""Word- und PDF-Dateien in Textzeilen mit Überschrift oder Seitenangabe zerlegen.

Grundlage ist read_docx/read_pdf aus app.py. Ergänzt: Word-Tabellen werden mitgelesen,
und PDF-Seiten ohne Textebene (Scans) werden als Lesehinweis gemeldet statt still übergangen.
"""
from __future__ import annotations

from io import BytesIO
from typing import IO

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

try:
    from PyPDF2 import PdfReader

    HAS_PYPDF2 = True
except Exception:  # pragma: no cover
    HAS_PYPDF2 = False

try:
    import pdfplumber

    HAS_PDFPLUMBER = True
except Exception:  # pragma: no cover
    HAS_PDFPLUMBER = False


class LeseFehler(Exception):
    """Die Datei konnte nicht gelesen werden oder enthält keinen Text.

    schluessel benennt die Meldung sprachneutral (siehe app/meldungen.py), die Oberfläche
    bekommt sie in ihrer Sprache; der Text hier bleibt Deutsch für Protokoll und Kommandozeile.
    """

    def __init__(self, meldung: str, schluessel: str | None = None) -> None:
        super().__init__(meldung)
        self.schluessel = schluessel


Struktur = list[dict[str, str]]
Hinweis = dict[str, object]  # {"art": "scan_seiten", "seiten": [2, 5], "gesamt": 7}


def _bytesio(daten: bytes | IO[bytes]) -> IO[bytes]:
    if isinstance(daten, (bytes, bytearray)):
        return BytesIO(daten)
    daten.seek(0)
    return daten


def docx_lesen(daten: bytes | IO[bytes]) -> Struktur:
    try:
        doc = Document(_bytesio(daten))
    except Exception as e:
        raise LeseFehler("Die Word-Datei lässt sich nicht öffnen. Vermutlich beschädigt oder kein echtes .docx.", "lesen_docx_oeffnen") from e
    structured: Struktur = []
    current_heading = "Unbekannt"

    def tabelle_lesen(table: Table, heading: str) -> None:
        for row in table.rows:
            gesehen: set[int] = set()
            for cell in row.cells:
                # Verbundene Zellen tauchen in row.cells mehrfach auf.
                if id(cell._tc) in gesehen:
                    continue
                gesehen.add(id(cell._tc))
                text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                if text:
                    structured.append({"text": text, "heading": f"{heading} (Tabelle)"})
                for innere in cell.tables:
                    tabelle_lesen(innere, heading)

    try:
        for child in doc.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                para = Paragraph(child, doc)
                text = para.text.strip()
                if not text:
                    continue
                style = getattr(para, "style", None)
                name = getattr(style, "name", "") if style is not None else ""
                if name.startswith("Heading"):
                    current_heading = text
                structured.append({"text": text, "heading": current_heading})
            elif tag == "tbl":
                tabelle_lesen(Table(child, doc), current_heading)
    except Exception as e:
        raise LeseFehler("Der Inhalt der Word-Datei lässt sich nicht lesen. Vermutlich beschädigt.", "lesen_docx_inhalt") from e
    if not structured:
        raise LeseFehler("Keine Textinhalte erkannt.", "lesen_kein_text")
    return structured


def _seitentexte_pypdf2(f: IO[bytes]) -> list[str]:
    reader = PdfReader(f)
    return [page.extract_text() or "" for page in reader.pages]


def _seitentexte_pdfplumber(f: IO[bytes]) -> list[str]:
    f.seek(0)
    with pdfplumber.open(f) as pdf:
        return [page.extract_text() or "" for page in pdf.pages]


def pdf_lesen_mit_hinweisen(daten: bytes | IO[bytes]) -> tuple[Struktur, list[Hinweis]]:
    """Text je Seite; Seiten ohne Textebene werden gesammelt und als Hinweis gemeldet."""
    f = _bytesio(daten)
    seiten: list[str] = []
    fehler: Exception | None = None
    if HAS_PYPDF2:
        try:
            seiten = _seitentexte_pypdf2(f)
        except Exception as e:
            fehler = e
    if HAS_PDFPLUMBER and (not seiten or not any(s.strip() for s in seiten)):
        try:
            seiten = _seitentexte_pdfplumber(f)
        except Exception as e:
            raise LeseFehler("Die PDF-Datei lässt sich nicht öffnen. Vermutlich beschädigt oder verschlüsselt.", "lesen_pdf_oeffnen") from e
    elif HAS_PDFPLUMBER and any(not s.strip() for s in seiten):
        # Gemischte PDFs: leere Seiten noch einmal mit dem zweiten Leser versuchen.
        try:
            ersatz = _seitentexte_pdfplumber(f)
            seiten = [s if s.strip() else (ersatz[i] if i < len(ersatz) else "") for i, s in enumerate(seiten)]
        except Exception:
            pass
    if not seiten and fehler is not None:
        raise LeseFehler("Die PDF-Datei lässt sich nicht öffnen. Vermutlich beschädigt oder verschlüsselt.", "lesen_pdf_oeffnen") from fehler
    structured: Struktur = []
    leer: list[int] = []
    for i, text in enumerate(seiten, start=1):
        zeilen = [z.strip() for z in text.splitlines() if z.strip()]
        if not zeilen:
            leer.append(i)
            continue
        structured.extend({"text": z, "heading": f"Seite {i}"} for z in zeilen)
    if not structured:
        raise LeseFehler("Keine Textinhalte erkannt. Vermutlich ein Scan ohne Textebene.", "lesen_kein_text_scan")
    hinweise: list[Hinweis] = []
    if leer:
        hinweise.append({"art": "scan_seiten", "seiten": leer, "gesamt": len(seiten)})
    return structured, hinweise


def pdf_lesen(daten: bytes | IO[bytes]) -> Struktur:
    return pdf_lesen_mit_hinweisen(daten)[0]


def lesen_mit_hinweisen(dateiname: str, daten: bytes | IO[bytes]) -> tuple[Struktur, list[Hinweis]]:
    """Struktur plus Lesehinweise (strukturiert, Text je Sprache über texte.lesehinweis)."""
    name = dateiname.lower()
    if name.endswith(".pdf"):
        return pdf_lesen_mit_hinweisen(daten)
    if name.endswith(".docx"):
        return docx_lesen(daten), []
    raise LeseFehler("Nur Word (.docx) und PDF (.pdf) werden geprüft.", "lesen_dateityp")


def lesen(dateiname: str, daten: bytes | IO[bytes]) -> Struktur:
    return lesen_mit_hinweisen(dateiname, daten)[0]
