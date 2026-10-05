"""Prüfung eines erzeugten Word-Dokuments, PDFs in mehreren Sprachen, ZIP, Ablegen, Mail-Entwurf."""
import io
import json
import zipfile

import pytest
from docx import Document
from PyPDF2 import PdfReader

from app.pruefer import berichte, pruefung, texte


def _docx(absaetze: list[str]) -> bytes:
    doc = Document()
    doc.add_heading("Betriebsanleitung Testgerät", level=1)
    for a in absaetze:
        doc.add_paragraph(a)
    puffer = io.BytesIO()
    doc.save(puffer)
    return puffer.getvalue()


LUECKENHAFT = [
    "Diese Anleitung richtet sich an die Zielgruppe der Benutzer.",
    "Ein Inhaltsverzeichnis ist am Anfang enthalten.",
    "Dieser Satz ist absichtlich sehr lang gehalten, damit die Prüfung auf überlange Sätze anspringt und wir sehen können, ob die Zusammenfassung der Fundstellen wie vorgesehen funktioniert, was bei kurzen Sätzen natürlich nicht passieren würde.",
    "Noch ein zweiter, ebenfalls viel zu langer Satz folgt hier, der ebenfalls mehr als fünfundzwanzig Wörter enthält und damit die Regel für die Satzlänge ein zweites Mal auslösen sollte, sodass die Zusammenfassung greift.",
]


def test_pruefen_direkt():
    struktur = [{"text": t, "heading": ""} for t in LUECKENHAFT]
    erg = pruefung.pruefen(struktur, ["din"])
    ids = {f["ID"] for f in erg["funde"]}
    assert "DIN-001" not in ids and "DIN-002" not in ids
    assert "DIN-003" in ids
    assert erg["score"] == 100 - sum(f.get("Gewichtung", 0) for f in erg["funde"])
    assert erg["ampel"] in ("gruen", "gelb", "rot")
    assert erg["fazit"]


def test_pruefung_hochladen_mit_zusatzsprache(client, admin):
    r = client.post(
        "/api/pruefung",
        files={"datei": ("Testgeraet_Anleitung.docx", _docx(LUECKENHAFT), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"regelsaetze": "basis,din,ce", "zusatzsprache": "en"},
    )
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["sprachen"] == ["de", "en"]
    assert len(d["pdfs"]) == 4
    assert {p["bericht"] for p in d["pdfs"]} == {"pruef", "fach"}
    assert d["pdfs"][0]["dateiname"].endswith("_Pruefbericht_DE_v01.pdf")
    assert d["fundeAnzahl"] >= 1
    assert d["regelnVorhanden"]["din"] is True
    pytest.pruef_id = d["id"]

    for p in d["pdfs"]:
        pr = client.get(p["url"])
        assert pr.status_code == 200, p
        assert pr.headers["content-type"].startswith("application/pdf")
        assert pr.content[:5] == b"%PDF-"
        assert len(pr.content) > 2000

    r = client.get(f"/api/pruefung/{d['id']}/pdf/pruef/ru")
    assert r.status_code == 404

    z = client.get(d["zip"])
    assert z.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(z.content))
    assert len([n for n in zf.namelist() if n.endswith(".pdf")]) == 4

    r = client.get("/api/pruefungen")
    assert any(p["id"] == d["id"] for p in r.json()["pruefungen"])
    r = client.get(f"/api/pruefung/{d['id']}")
    assert r.status_code == 200 and r.json()["id"] == d["id"]


def test_pruefung_ablegen_und_mailentwurf(client, admin):
    pid = pytest.pruef_id
    r = client.post(f"/api/pruefung/{pid}/ablegen")
    assert r.status_code == 200, r.text
    assert len(r.json()["abgelegt"]) == 4
    r = client.post(f"/api/pruefung/{pid}/mail-entwurf", json={"an": "kunde@example.com"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["entwurf"] in ("text", "mail")
    assert "Fachbericht" in d["betreff"] or "Testgeraet" in d["betreff"]
    assert len(d["anhaenge"]) == 2
    assert all("Fachbericht" in a for a in d["anhaenge"])
    if d["entwurf"] == "text":
        assert "Ukraine" not in d["text"]


def test_pruefung_in_ukrainisch_und_russisch(client, admin):
    client.patch("/api/ich/einstellungen", json={"language": "uk"})
    r = client.post(
        "/api/pruefung",
        files={"datei": ("Pruefung.docx", _docx(LUECKENHAFT[:1]), "application/octet-stream")},
        data={"regelsaetze": "din", "zusatzsprache": "ru"},
    )
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["sprachen"] == ["uk", "ru"]
    for p in d["pdfs"]:
        pr = client.get(p["url"])
        assert pr.status_code == 200 and pr.content[:5] == b"%PDF-"
    client.patch("/api/ich/einstellungen", json={"language": "de"})


def test_pruefung_lehnt_falsche_dateien_ab(client, admin):
    r = client.post("/api/pruefung", files={"datei": ("x.txt", b"hallo", "text/plain")}, data={})
    assert r.status_code == 400
    r = client.post("/api/pruefung", files={"datei": ("kaputt.docx", b"kein zip", "application/octet-stream")}, data={})
    assert r.status_code == 422
    r = client.post("/api/pruefung", files={"datei": ("x.docx", _docx(["a"]), "application/octet-stream")}, data={"regelsaetze": "nix"})
    assert r.status_code == 400


def test_fremde_pruefung_nicht_sichtbar(admin):
    from fastapi.testclient import TestClient
    from app import auth, db
    from app.main import app
    from tests.conftest import PASSWORT

    with TestClient(app) as c:
        with db.transaktion() as con:
            auth.benutzer_anlegen(con, "fremd@test.local", "Fremd", "mitglied")
            con.execute("UPDATE users SET status='aktiv', passwort_hash=? WHERE email='fremd@test.local'", (auth.hash_passwort(PASSWORT),))
        assert c.post("/api/anmelden", json={"email": "fremd@test.local", "passwort": PASSWORT}).status_code == 200
        assert c.get(f"/api/pruefung/{pytest.pruef_id}").status_code == 404
        assert c.get("/api/pruefungen").json()["pruefungen"] == []


def test_berichte_direkt_alle_sprachen():
    struktur = [{"text": LUECKENHAFT[0], "heading": ""}]
    erg = pruefung.pruefen(struktur, ["din"])
    erg["pruefer"] = "Test"
    meta = {"dateiname": "Doku.docx", "erstellt": "2026-10-01T09:00:00", "pruefer": "Test"}
    for sp in ("de", "en", "uk", "ru"):
        for art in ("pruef", "fach"):
            pdf = berichte.erzeugen(art, erg, meta, sp)
            assert pdf[:5] == b"%PDF-", (art, sp)
            # Der Fuß trägt nur die fertige Seitenzahl, keinen Platzhalter aus dem Zähldurchlauf.
            text = " ".join(" ".join(s.extract_text().split()) for s in PdfReader(io.BytesIO(pdf)).pages)
            assert "{m}" not in text, (art, sp)
            assert texte.t(sp, "seite", n=1, m=len(PdfReader(io.BytesIO(pdf)).pages)) in text, (art, sp)
    assert berichte.dateiname("fach", "Doku.docx", "uk", "2026-10-01T09:00:00") == "2026-10-01_Doku_Fachbericht_UK_v01.pdf"
    assert texte.t("uk", "pruefbericht")


def test_pruefung_mit_fortschritt(client, admin):
    """Der Zeilenstrom meldet die Phasen je Dokument, ein Ergebnis je Dokument und zuletzt den Lauf."""
    with client.stream(
        "POST",
        "/api/pruefung",
        files={"datei": ("Strom.docx", _docx(LUECKENHAFT[:2]), "application/octet-stream")},
        data={"regelsaetze": "din", "zusatzsprache": "en", "fortschritt": "1"},
    ) as r:
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/x-ndjson")
        zeilen = [json.loads(z) for z in r.iter_lines() if z.strip()]
    phasen = [z.get("phase") for z in zeilen]
    assert phasen[:2] == ["lesen", "pruefen"]
    assert phasen.count("berichte") == 4
    assert [z for z in zeilen if z.get("phase") == "berichte"][-1] == {"phase": "berichte", "n": 4, "von": 4, "dokument": 1, "dokumente": 1, "datei": "Strom.docx"}
    assert phasen[-2:] == ["ergebnis", "fertig"]
    assert zeilen[-2]["ergebnis"]["id"] == zeilen[-1]["lauf"]["ergebnisse"][0]["id"]
    erg = zeilen[-1]["lauf"]["ergebnisse"][0]
    assert len(erg["pdfs"]) == 4 and erg["fundeAnzahl"] >= 1

    # Eine unlesbare Datei ist ein Ergebnis mit Fehler, der Lauf endet trotzdem ordentlich.
    with client.stream(
        "POST",
        "/api/pruefung",
        files={"datei": ("kaputt.docx", b"kein zip", "application/octet-stream")},
        data={"fortschritt": "1"},
    ) as r:
        zeilen = [json.loads(z) for z in r.iter_lines() if z.strip()]
    assert zeilen[-2]["phase"] == "ergebnis" and zeilen[-2]["status"] == 422 and "fehler" in zeilen[-2]
    assert zeilen[-1]["phase"] == "fertig" and zeilen[-1]["lauf"]["ergebnisse"] == [] and len(zeilen[-1]["lauf"]["fehler"]) == 1


def test_mehrere_dateien_in_einem_lauf(client, admin, tmp_path, monkeypatch):
    """Bis zu 20 Dateien je Prüfung: jede bekommt ihre Berichte, eine kaputte stoppt die anderen nicht."""
    from app import config

    monkeypatch.setattr(config, "OUTPUT", tmp_path / "output")
    dateien = [
        ("datei", ("Eins.docx", _docx(LUECKENHAFT), "application/octet-stream")),
        ("datei", ("Zwei.docx", _docx(LUECKENHAFT[:1]), "application/octet-stream")),
        ("datei", ("Kaputt.docx", b"kein zip", "application/octet-stream")),
    ]
    r = client.post("/api/pruefung", files=dateien, data={"regelsaetze": "basis,din", "zusatzsprache": "en"})
    assert r.status_code == 200, r.text
    lauf = r.json()
    assert lauf["dateien"] == 3 and len(lauf["ergebnisse"]) == 2 and len(lauf["fehler"]) == 1
    assert lauf["fehler"][0]["datei"] == "Kaputt.docx" and lauf["fehler"][0]["status"] == 422
    assert [e["dateiname"] for e in lauf["ergebnisse"]] == ["Eins.docx", "Zwei.docx"]
    assert lauf["pdfAnzahl"] == 8
    assert all(e["sprachen"] == ["de", "en"] for e in lauf["ergebnisse"])

    r = client.get(f"/api/lauf/{lauf['lauf']}")
    assert r.status_code == 200 and len(r.json()["ergebnisse"]) == 2
    z = client.get(lauf["zip"])
    assert z.status_code == 200
    namen = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
    assert len(namen) == 8 and any("Eins_" in n for n in namen) and any("Zwei_" in n for n in namen)
    r = client.post(lauf["ablegen"])
    assert r.status_code == 200 and len(r.json()["abgelegt"]) == 8 and r.json()["dokumente"] == 2

    liste = client.get("/api/pruefungen").json()["pruefungen"]
    assert [p["lauf"] for p in liste[:2]] == [lauf["lauf"], lauf["lauf"]]
    assert client.get(f"/api/lauf/{lauf['lauf']}x").status_code == 404

    # Mehrfachlauf mit Fortschritt: Dokumentzähler und ein Ergebnis je Datei.
    with client.stream("POST", "/api/pruefung", files=dateien[:2], data={"regelsaetze": "din", "fortschritt": "1"}) as r:
        zeilen = [json.loads(x) for x in r.iter_lines() if x.strip()]
    ergebnisse = [x for x in zeilen if x.get("phase") == "ergebnis"]
    assert [(x["dokument"], x["dokumente"], x["datei"]) for x in ergebnisse] == [(1, 2, "Eins.docx"), (2, 2, "Zwei.docx")]
    assert zeilen[0] == {"phase": "lesen", "dokument": 1, "dokumente": 2, "datei": "Eins.docx"}
    assert zeilen[-1]["phase"] == "fertig" and len(zeilen[-1]["lauf"]["ergebnisse"]) == 2

    zu_viele = [("datei", (f"D{i}.docx", _docx(["a"]), "application/octet-stream")) for i in range(config.MAX_DATEIEN + 1)]
    r = client.post("/api/pruefung", files=zu_viele, data={"regelsaetze": "din"})
    assert r.status_code == 400 and str(config.MAX_DATEIEN) in r.json()["fehler"]
    r = client.post("/api/pruefung", files=dateien[:1] + [("datei", ("x.txt", b"hallo", "text/plain"))], data={})
    assert r.status_code == 400 and r.json()["fehler"].startswith("x.txt")


def test_word_tabellen_werden_gelesen():
    from app.pruefer import lesen
    doc = Document()
    doc.add_paragraph("Einleitung ohne Stichwörter.")
    tab = doc.add_table(rows=2, cols=2)
    tab.cell(0, 0).text = "Sicherheitshinweise"
    tab.cell(0, 1).text = "Vor der Montage lesen."
    tab.cell(1, 0).text = "Garantie"
    tab.cell(1, 1).text = "24 Monate Gewährleistung."
    puffer = io.BytesIO()
    doc.save(puffer)
    struktur = lesen.docx_lesen(puffer.getvalue())
    texte_ = [z["text"] for z in struktur]
    assert "Sicherheitshinweise" in texte_ and "24 Monate Gewährleistung." in texte_
    assert any(z["heading"].endswith("(Tabelle)") for z in struktur)
    erg = pruefung.pruefen(struktur, ["basis"])
    ids = {f["ID"] for f in erg["funde"]}
    assert "CHK-002" not in ids and "CHK-008" not in ids


def test_pdf_scanseiten_werden_gemeldet(tmp_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from app.pruefer import lesen

    pfad = tmp_path / "gemischt.pdf"
    c = canvas.Canvas(str(pfad), pagesize=A4)
    c.drawString(72, 750, "Diese Anleitung richtet sich an die Zielgruppe der Benutzer.")
    c.showPage()
    c.rect(100, 100, 200, 200)  # Seite 2: nur Grafik, keine Textebene
    c.showPage()
    c.drawString(72, 750, "Ein Inhaltsverzeichnis ist enthalten.")
    c.showPage()
    c.save()
    struktur, hinweise = lesen.lesen_mit_hinweisen("gemischt.pdf", pfad.read_bytes())
    assert len(struktur) == 2
    assert hinweise == [{"art": "scan_seiten", "seiten": [2], "gesamt": 3}]
    assert texte.lesehinweis("de", hinweise[0]) == "Seite 2 von 3 Seiten hat keine Textebene und wurde nicht geprüft. Vermutlich gescannt."
    assert texte.lesehinweis("en", {"art": "scan_seiten", "seiten": [2, 5], "gesamt": 7}).startswith("Pages 2, 5 of 7 pages")

    nur_bild = tmp_path / "scan.pdf"
    c = canvas.Canvas(str(nur_bild), pagesize=A4)
    c.rect(100, 100, 200, 200)
    c.showPage()
    c.save()
    with pytest.raises(lesen.LeseFehler, match="Scan"):
        lesen.lesen_mit_hinweisen("scan.pdf", nur_bild.read_bytes())


def test_ablegen_ueberschreibt_nicht(client, admin, tmp_path, monkeypatch):
    from app import config
    monkeypatch.setattr(config, "OUTPUT", tmp_path)
    r = client.post(
        "/api/pruefung",
        files={"datei": ("Ablage.docx", _docx(LUECKENHAFT[:1]), "application/octet-stream")},
        data={"regelsaetze": "din", "zusatzsprache": ""},
    )
    pid = r.json()["id"]
    erste = client.post(f"/api/pruefung/{pid}/ablegen").json()["abgelegt"]
    zweite = client.post(f"/api/pruefung/{pid}/ablegen").json()["abgelegt"]
    assert erste == zweite and all(n.endswith("_v01.pdf") for n in erste)
    # Fremde Datei gleichen Namens mit anderem Inhalt darf nicht überschrieben werden.
    fremd = tmp_path / erste[0].split("/")[-1]
    fremd.write_bytes(b"%PDF-fremd")
    dritte = client.post(f"/api/pruefung/{pid}/ablegen").json()["abgelegt"]
    assert fremd.read_bytes() == b"%PDF-fremd"
    assert dritte[0].endswith("_v02.pdf")


def test_pruefpunkte_auswaehlen(client, admin):
    """Einzelne Prüfpunkte lassen sich auslassen; das steht im Ergebnis und im Bericht, der Score zählt sie nicht."""
    r = client.get("/api/regeln")
    assert r.status_code == 200, r.text
    saetze = r.json()["regelsaetze"]
    assert set(saetze) == {"basis", "din", "ce"} and r.json()["ausgelassen"] == []
    basis_ids = [p["id"] for p in saetze["basis"]["punkte"]]
    assert "TXT-001" in basis_ids and "CHK-008" in basis_ids and len(basis_ids) == 9
    assert len(saetze["din"]["punkte"]) == 8 and len(saetze["ce"]["punkte"]) == 8

    r = client.put("/api/ich/pruefpunkte", json={"ausgelassen": ["DIN-008", "CHK-008", "TXT-001"]})
    assert r.status_code == 200 and r.json()["ausgelassen"] == ["CHK-008", "DIN-008", "TXT-001"]
    assert client.get("/api/regeln").json()["ausgelassen"] == ["CHK-008", "DIN-008", "TXT-001"]
    assert client.put("/api/ich/pruefpunkte", json={"ausgelassen": ["NIX-1"]}).status_code == 400
    assert client.put("/api/ich/pruefpunkte", json={"ausgelassen": "CHK-001"}).status_code == 400

    datei = {"datei": ("Punkte.docx", _docx(LUECKENHAFT), "application/octet-stream")}
    voll = client.post("/api/pruefung", files=datei, data={"regelsaetze": "basis,din,ce"}).json()
    r = client.post("/api/pruefung", files=datei, data={"regelsaetze": "basis,din,ce", "zusatzsprache": "en", "ausgelassen": "CHK-008,DIN-008,TXT-001"})
    assert r.status_code == 200, r.text
    d = r.json()
    ids = {f["ID"] for f in d["funde"]}
    assert not ids & {"CHK-008", "DIN-008", "TXT-001"}
    assert {f["ID"] for f in voll["funde"]} >= {"CHK-008", "DIN-008", "TXT-001"}
    assert [a["id"] for a in d["ausgelassen"]] == ["CHK-008", "TXT-001", "DIN-008"]
    assert d["punkteGesamt"] == 25 and d["punkteGeprueft"] == 22
    assert d["score"] == 100 - sum(f["Gewichtung"] for f in d["funde"]) and d["score"] > voll["score"]
    for sp, erwartet in (("de", "22 von 25 geprüft. Bewusst ausgelassen: CHK-008 Garantie, TXT-001 Lesbarkeit, DIN-008 Format"), ("en", "22 of 25 checked. Deliberately left out: CHK-008 ")):
        pdf = client.get(f"/api/pruefung/{d['id']}/pdf/pruef/{sp}").content
        text = " ".join(" ".join(s.extract_text().split()) for s in PdfReader(io.BytesIO(pdf)).pages)
        assert erwartet in text, (sp, text[:600])

    # Unbekannter Punkt und leerer Regelsatz sind Fehler, bevor etwas geprüft wird
    r = client.post("/api/pruefung", files=datei, data={"regelsaetze": "din", "ausgelassen": "CHK-001"})
    assert r.status_code == 400 and "CHK-001" in r.json()["fehler"]
    alle_din = ",".join(p["id"] for p in saetze["din"]["punkte"])
    r = client.post("/api/pruefung", files=datei, data={"regelsaetze": "basis,din", "ausgelassen": alle_din})
    assert r.status_code == 400 and "din" in r.json()["fehler"]
    client.put("/api/ich/pruefpunkte", json={"ausgelassen": []})


def test_fremdsprachige_berichte_ohne_deutsche_regeltexte():
    """In en, uk und ru steht kein deutscher Regeltext mehr: Bereiche und Empfehlungen aller drei Regelsätze sind übersetzt."""
    import json

    from app import config

    struktur = [{"text": LUECKENHAFT[0], "heading": ""}]
    erg = pruefung.pruefen(struktur, ["basis", "din", "ce"], ["CHK-002"])
    erg["pruefer"] = "Test"
    assert len(erg["funde"]) >= 20
    meta = {"dateiname": "Doku.docx", "erstellt": "2026-10-05T09:00:00", "pruefer": "Test"}
    regeln_alle = [r for datei in ("pruefkatalog.json", "normlogik_82079.json", "ce_logik.json") for r in json.loads((config.REGELN / datei).read_text(encoding="utf-8"))]
    for sp in ("en", "uk", "ru"):
        text = ""
        for art in ("pruef", "fach"):
            pdf = berichte.erzeugen(art, erg, meta, sp)
            text += " " + " ".join(" ".join(s.extract_text().split()) for s in PdfReader(io.BytesIO(pdf)).pages)
        eng = "".join(text.split())  # Zeilenumbrüche und Trennstriche stören den Vergleich nicht
        funde_ids = {f["ID"] for f in erg["funde"]}
        for r in regeln_alle:
            if r["id"] in funde_ids:
                assert "".join(r[f"empfehlung_{sp}"].split()) in eng, (sp, r["id"], "Übersetzung fehlt im Bericht")
            assert "".join(r["empfehlung"].split()) not in eng, (sp, r["id"], "deutsche Empfehlung im Bericht")
            if r[f"bereich_{sp}"] != r["bereich"]:
                assert r["bereich"] not in text, (sp, r["id"], "deutscher Bereich im Bericht")
        assert "Sicherheit" not in text and "Garantie" not in text and "nachweisen" not in text, sp
