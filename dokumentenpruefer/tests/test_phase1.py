"""Abnahme-Gegenproben für lokale Prüfung, Berichte und Betrieb."""
import io
from concurrent.futures import ThreadPoolExecutor

import pytest
from docx import Document
from PyPDF2 import PdfReader

from app import config, db
from app.pruefer import berichte, lesen, pruefung, texte


def dokument():
    doc = Document()
    doc.add_heading("Testanleitung", 1)
    doc.add_paragraph("Kurzer Text zur Installation.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("strom", [False, True])
def test_berichtfehler_hinterlaesst_keinen_erfolg(client, admin, monkeypatch, strom):
    with db.transaktion() as con:
        vorher = con.execute("SELECT COUNT(*) FROM pruefungen").fetchone()[0]
    ordner = set(config.PRUEFUNGEN.iterdir())
    original = berichte.erzeugen
    aufrufe = 0

    def kaputt(*args):
        nonlocal aufrufe
        aufrufe += 1
        if aufrufe == 2:
            raise OSError("Simulierter Schreibfehler")
        return original(*args)

    monkeypatch.setattr(berichte, "erzeugen", kaputt)
    r = client.post("/api/pruefung", files={"datei": ("Fehler.docx", dokument())},
                    data={"regelsaetze": "din", "fortschritt": "1" if strom else ""})
    if strom:
        import json
        ereignisse = [json.loads(z) for z in r.text.splitlines()]
        assert ereignisse[-1]["status"] == 500
        assert not any(e.get("phase") == "fertig" for e in ereignisse)
    else:
        assert r.status_code == 500
        assert "fehler" in r.json()
    with db.transaktion() as con:
        assert con.execute("SELECT COUNT(*) FROM pruefungen").fetchone()[0] == vorher
    assert set(config.PRUEFUNGEN.iterdir()) == ordner


def test_parallele_ablage_ueberschreibt_nicht(tmp_path):
    inhalte = [b"%PDF-Probe-" + str(i).encode() for i in range(12)]
    def ablegen(inhalt):
        return berichte.ablegen(tmp_path, "pruef", "Gleich.docx", "de", "2026-10-02", inhalt)
    with ThreadPoolExecutor(max_workers=12) as pool:
        pfade = list(pool.map(ablegen, inhalte))
    assert len(set(pfade)) == 12
    assert [p.read_bytes() for p in pfade] == inhalte
    assert ablegen(inhalte[0]) == pfade[0]


def test_verschachtelte_tabellen_behalten_reihenfolge():
    doc = Document()
    doc.add_heading("Wartung", 1)
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Vorher"
    inner = cell.add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "Innen"
    cell.add_paragraph("Nachher")
    buf = io.BytesIO(); doc.save(buf)
    struktur = lesen.docx_lesen(buf.getvalue())
    assert [s["text"] for s in struktur] == ["Wartung", "Vorher", "Innen", "Nachher"]


def test_vertikal_verbundene_zelle_nur_einmal():
    doc = Document()
    tab = doc.add_table(rows=2, cols=2)
    tab.cell(0, 0).merge(tab.cell(1, 0)).text = "Einmaliger Hinweis"
    tab.cell(0, 1).text = "Erste Zeile"
    tab.cell(1, 1).text = "Zweite Zeile"
    buf = io.BytesIO(); doc.save(buf)
    assert [s["text"] for s in lesen.docx_lesen(buf.getvalue())].count("Einmaliger Hinweis") == 1


def test_pruefbericht_enthaelt_fundstelle():
    erg = pruefung.pruefen([{"text": " ".join(["Wort"] * 30), "heading": "Wartung / Seite 17"}], ["basis"])
    pdf = berichte.erzeugen("pruef", erg, {"dateiname": "Test.docx"}, "de")
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "Wartung / Seite 17" in text


@pytest.mark.parametrize("sprache", ["en", "uk", "ru"])
def test_alle_regeln_haben_berichtuebersetzungen(sprache):
    from app.pruefer import regeln
    for key in config.REGELSAETZE:
        for regel in regeln.laden(key):
            assert regel.get("empfehlung_" + sprache), (key, regel["id"], sprache)


def test_ungueltige_zusatzsprache_wird_abgewiesen(client, admin):
    r = client.post("/api/pruefung", files={"datei": ("Probe.docx", dokument())},
                    data={"zusatzsprache": "zz"})
    assert r.status_code == 400
