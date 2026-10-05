"""Gegenproben für Inhaltsnachweise und PDF-Zeilen, ausschließlich Testmaterial."""
import io

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from app.pruefer import lesen, pruefung, regeln


def word(*abschnitte):
    d = Document()
    for art, text in abschnitte:
        if art == 'titel':
            d.add_heading(text, 1)
        else:
            d.add_paragraph(text)
    b = io.BytesIO(); d.save(b)
    return lesen.docx_lesen(b.getvalue())


def hat_fund(struktur, regel='CHK-002', regelsatz='basis'):
    e = pruefung.pruefen(struktur, [regelsatz])
    fund = any(f['ID'] == regel for f in e['funde'])
    suchstatus = next(r['suchstatus'] for r in e['regelpruefungen'] if r['id'] == regel)
    assert (suchstatus == 'nicht_gefunden') == fund
    assert e['freigabe'] is False
    return fund


def test_leere_word_kapitel_sind_keine_inhaltsnachweise():
    struktur = word(*(('titel', r['keywords'][0]) for r in regeln.laden('basis')))
    e = pruefung.pruefen(struktur, ['basis'])
    assert {f['ID'] for f in e['funde']} == {r['id'] for r in regeln.laden('basis')}
    assert e['suchtrefferAnzahl'] == 0
    assert e['score'] == 76 and e['ampel'] == 'gelb'  # Bestehende Rechnung unverändert.


def test_word_titel_mit_konkretem_inhalt_zaehlt():
    assert not hat_fund(word(('titel', 'Sicherheitshinweise'),
                            ('absatz', 'Vor Arbeiten den Hauptschalter ausschalten und gegen Wiedereinschalten sichern.')))


def test_leeres_kapitel_bekommt_nicht_den_inhalt_des_naechsten():
    assert hat_fund(word(('titel', 'Sicherheitshinweise'), ('titel', 'Wartung'),
                        ('absatz', 'Den Filter nach 100 Betriebsstunden wechseln.')))


@pytest.mark.parametrize('text', ['TODO', 'Folgt.', 'Wird ergänzt.', 'Noch zu ergänzen.', 'Nicht vorhanden.'])
def test_platzhalter_unter_ueberschrift_ist_kein_nachweis(text):
    assert hat_fund(word(('titel', 'Sicherheitshinweise'), ('absatz', text)))


@pytest.mark.parametrize('text', [
    'Sicherheitshinweise fehlen vollständig.',
    'Sicherheitshinweise sind nicht vorhanden.',
    'Es fehlen die Sicherheitshinweise.',
    'Diese Anleitung enthält keine Sicherheitshinweise.',
    'Die Sicherheitshinweise wurden nicht beschrieben.',
    'Sicherheitshinweise: noch zu ergänzen.',
    'Die Dokumentation liegt ohne Sicherheitshinweise vor.',
])
def test_aussage_ueber_fehlende_sicherheit_ist_kein_nachweis(text):
    assert hat_fund(word(('absatz', text)))


@pytest.mark.parametrize('text', [
    'Vorsicht: Die heißen Bauteile nicht berühren.',
    'Sicherheitshinweise beachten. Den Motor niemals mit geöffneter Abdeckung starten.',
    'Bei Nichtbeachtung der Sicherheitshinweise können Verletzungen entstehen.',
    'Die Montageanleitung fehlt, aber die Sicherheitshinweise sind vollständig enthalten.',
])
def test_verbot_oder_andere_fehlende_angabe_entwertet_keinen_sicherheitshinweis(text):
    assert not hat_fund(word(('absatz', text)))


@pytest.mark.parametrize('regel,regelsatz,text', [
    ('DIN-002', 'din', 'Ein Inhaltsverzeichnis fehlt.'),
    ('CE-003', 'ce', 'Herstellerangaben sind nicht vorhanden.'),
])
def test_absagen_gelten_in_allen_regelsaetzen(regel, regelsatz, text):
    assert hat_fund(word(('absatz', text)), regel, regelsatz)


def test_negierter_treffer_verdeckt_keinen_spaeteren_hinweis():
    assert not hat_fund(word(('absatz', 'Sicherheitshinweise fehlen im Entwurf.'),
                            ('absatz', 'Vorsicht: Vor dem Öffnen den Netzstecker ziehen.')))


def test_nachweis_wird_nicht_ueber_zwei_word_abschnitte_erfunden():
    assert hat_fund(word(('absatz', 'technische'), ('absatz', 'Daten')), 'CHK-006')


def test_tabellenbeschriftung_braucht_inhalt_in_derselben_zeile():
    d = Document(); t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = 'Sicherheitshinweise'
    t.cell(1, 0).text = 'Wartung'; t.cell(1, 1).text = 'Filter täglich wechseln.'
    b = io.BytesIO(); d.save(b)
    assert hat_fund(lesen.docx_lesen(b.getvalue()))
    t.cell(0, 1).text = 'Vor Arbeiten die Maschine ausschalten.'
    b = io.BytesIO(); d.save(b)
    assert not hat_fund(lesen.docx_lesen(b.getvalue()))


def pdf(zeilen):
    b = io.BytesIO(); c = canvas.Canvas(b)
    for text, x, y, schrift, groesse in zeilen:
        c.setFont(schrift, groesse); c.drawString(x, y, text)
    c.save(); return b.getvalue()


def langfunde(struktur):
    return [f for f in pruefung.pruefen(struktur, ['basis'])['funde'] if f['ID'] == 'TXT-001']


def test_pdf_umbrueche_eines_satzes_werden_zusammengefuehrt():
    b = pdf([(' '.join(['Wort'] * 13), 72, 750, 'Helvetica', 11),
             (' '.join(['Wort'] * 14) + '.', 72, 736, 'Helvetica', 11)])
    funde = langfunde(lesen.pdf_lesen(b))
    assert len(funde) == 1 and funde[0]['Anzahl'] == 1
    assert 'Seite 1' in funde[0]['Fundstelle']


def test_pdf_trennt_ueberschrift_und_naechsten_absatz():
    b = pdf([(' '.join(['Titel'] * 13), 72, 750, 'Helvetica-Bold', 14),
             (' '.join(['Wort'] * 14) + '.', 72, 730, 'Helvetica', 11)])
    assert langfunde(lesen.pdf_lesen(b)) == []


def test_pdf_spalten_werden_nicht_zu_einem_satz_verkettet():
    b = pdf([(' '.join(['a'] * 13), 72, 750, 'Helvetica', 11),
             (' '.join(['b'] * 14) + '.', 320, 750, 'Helvetica', 11)])
    assert langfunde(lesen.pdf_lesen(b)) == []


def test_pdf_listenpunkte_bleiben_getrennt():
    b = pdf([('- ' + ' '.join(['a'] * 13), 72, 750, 'Helvetica', 11),
             ('- ' + ' '.join(['b'] * 14), 72, 736, 'Helvetica', 11)])
    assert langfunde(lesen.pdf_lesen(b)) == []


def test_pdf_toc_eintraege_ersetzen_keine_sicherheitshinweise():
    b = pdf([('Inhaltsverzeichnis', 72, 770, 'Helvetica-Bold', 16),
             ('Sicherheitshinweise .......................... 7', 72, 740, 'Helvetica', 11)])
    assert hat_fund(lesen.pdf_lesen(b))


def test_pdf_titel_mit_inhalt_bleibt_ein_suchtreffer():
    b = pdf([('Sicherheitshinweise', 72, 770, 'Helvetica-Bold', 16),
             ('Vor Arbeiten die Maschine ausschalten.', 72, 740, 'Helvetica', 11)])
    assert not hat_fund(lesen.pdf_lesen(b))


@pytest.mark.parametrize('abkuerzung', ['usw.', 'etc.'])
def test_abkuerzung_mit_kleingeschriebener_fortsetzung_bleibt_im_satz(abkuerzung):
    text = ' '.join(['Wort'] * 13) + ' ' + abkuerzung + ' ' + ' '.join(['weiter'] * 13) + '.'
    assert len(langfunde([{'text': text, 'heading': 'Wartung'}])) == 1


@pytest.mark.parametrize('abkuerzung', ['usw.', 'etc.'])
def test_abkuerzung_am_echten_satzende_bleibt_satzende(abkuerzung):
    a = ' '.join(['Wort'] * 12) + ' ' + abkuerzung
    b = 'Danach ' + ' '.join(['weiter'] * 13) + '.'
    assert langfunde([{'text': a + ' ' + b, 'heading': 'Wartung'}]) == []


def test_pdf_positionierte_woerter_behalten_wortabstaende():
    b = io.BytesIO(); c = canvas.Canvas(b); c.setFont('Helvetica', 11)
    for y, anzahl in [(750, 13), (736, 14)]:
        x = 72
        for i in range(anzahl):
            text = 'Wort.' if y == 736 and i == 13 else 'Wort'
            c.drawString(x, y, text)
            x += c.stringWidth(text, 'Helvetica', 11) + 2.5
    c.save()
    assert len(langfunde(lesen.pdf_lesen(b.getvalue()))) == 1
