"""Satzlänge innerhalb eines Textabschnitts statt Länge des ganzen Absatzes."""
import io
import subprocess
import sys

import pytest
from docx import Document

from app.pruefer import lesen, pruefung


KURZ = "Vor dem Start muss der Bediener alle vorhandenen Schutzeinrichtungen kontrollieren"


def funde(text):
    return pruefung.check_document([{"text": text, "heading": "Wartung"}], [])


@pytest.mark.parametrize("ende", [".", "!", "?", '.”', '."', ".»", ".)"])
def test_mehrere_kurze_saetze_sind_kein_langer_satz(ende):
    assert len(KURZ.split()) == 10
    assert funde(" ".join([KURZ + ende] * 3)) == []


@pytest.mark.parametrize("anzahl,erwartet", [(25, 0), (26, 1)])
def test_grenze_bleibt_bei_mehr_als_25_woertern(anzahl, erwartet):
    assert len(funde(" ".join(["Wort"] * anzahl) + ".")) == erwartet


def test_nur_der_lange_satz_wird_gemeldet():
    lang = " ".join(["Prüfwort"] * 26) + "."
    ergebnis = funde(KURZ + ". " + lang + " " + KURZ + ".")
    assert len(ergebnis) == 1
    assert ergebnis[0]["Empfehlung"] == lang[:150]
    assert ergebnis[0]["Fundstelle"] == "Wartung"


def test_zwei_lange_saetze_werden_als_zwei_gezaehlt():
    lang = " ".join(["Wort"] * 26) + "."
    ergebnis = pruefung.deduplicate_findings(funde(lang + " " + lang))
    assert len(ergebnis) == 1
    assert ergebnis[0]["Anzahl"] == 2
    assert ergebnis[0]["Gewichtung"] == 2  # Bestehende Sammelgewichtung bleibt erhalten.


@pytest.mark.parametrize("einlage", [
    "z. B.", "z.B.", "d. h.", "u. a.", "i. d. R.", "ca.", "bzw.",
    "Nr. 3", "Abb. 2", "Abs. 1", "Dr. Müller", "e. g.",
    "1.5 bar", "Version 1.2.3", "am 3. Oktober", "im 2. Schritt",
])
def test_abkuerzungen_und_zahlen_zerlegen_keinen_langen_satz(einlage):
    text = " ".join(["Wort"] * 13) + " " + einlage + " " + " ".join(["Wort"] * 13) + "."
    assert len(funde(text)) == 1


def test_zeilenumbruch_innerhalb_eines_abschnitts_beendet_keinen_satz():
    assert len(funde(" ".join(["Wort"] * 15) + "\n" + " ".join(["Wort"] * 15))) == 1


def test_leere_abschnitte_und_rest_ohne_schlusspunkt():
    assert funde(" \n ") == []
    assert funde(KURZ + ". " + KURZ + ". " + KURZ) == []
    assert len(funde(" ".join(["Wort"] * 26))) == 1


def test_word_absatz_und_tabellenzelle_mit_mehreren_saetzen():
    doc = Document()
    text = ". ".join([KURZ] * 3) + "."
    doc.add_paragraph(text)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = text
    puffer = io.BytesIO()
    doc.save(puffer)
    struktur = lesen.docx_lesen(puffer.getvalue())
    assert len(struktur) == 2
    ergebnis = pruefung.pruefen(struktur, ["basis"])
    assert not any(f["ID"] == "TXT-001" for f in ergebnis["funde"])


@pytest.mark.parametrize("folge", [".", "a."])
def test_lange_punktfolgen_blockieren_die_pruefung_nicht(folge):
    # Eigener Prozess: Ein Rückfall wird beendet und blockiert nicht den Testlauf.
    # 50.000 Wiederholungen dauern linear deutlich unter einer Sekunde; der alte
    # quadratische Ausdruck überschreitet das großzügige Zehn-Sekunden-Limit.
    code = (
        "import sys; from app.pruefer.pruefung import _saetze; "
        "text = sys.argv[1] * 50000 + 'X'; assert _saetze(text) == [text]"
    )
    subprocess.run([sys.executable, "-c", code, folge], check=True, timeout=10,
                   capture_output=True, text=True)


@pytest.mark.parametrize("text,erwartet", [
    ("z. B. Dieses Beispiel.", ["z. B. Dieses Beispiel."]),
    ("a. b. c.X", ["a. b. c.X"]),
    ("a.b.X", ["a.b.X"]),
    ('Erster Satz.!\" Zweiter Satz.', ['Erster Satz.!\"', 'Zweiter Satz.']),
])
def test_satzgrenzen_nach_punktfolgen_bleiben_erhalten(text, erwartet):
    assert pruefung._saetze(text) == erwartet
